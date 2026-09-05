# -*- coding: utf-8 -*-
"""Conclusions: add certainty grading + metabolomic closure (EN + CN), then
recompute EN word counts."""
import re
from docx import Document

BASE = "方案B_合并稿/投稿_HeartRhythm_压缩版"


def set_text(p, t):
    p.runs[0].text = t
    for r in p.runs[1:]:
        r.text = ""


def replace_in_doc(doc, old, new):
    for p in doc.paragraphs:
        if old in p.text:
            set_text(p, p.text.replace(old, new))
            return
    raise RuntimeError("not found: " + old[:50])


d = Document(f"{BASE}/咖啡与房颤_投稿版_v2.docx")
replace_in_doc(d,
    "objective biomarker MR was null; and no real-world safety signal emerged.",
    "objective biomarker MR was null; genetic intake liability did not raise "
    "circulating caffeine-family metabolites; and no real-world safety signal emerged.")
replace_in_doc(d,
    "These data argue against a moderate-to-large causal effect of customary coffee "
    "consumption on AF, although small effects in either direction cannot be excluded; "
    "routine caffeine restriction, including after cardioversion, remains unsupported.",
    "These data argue with moderate certainty against a moderate-to-large causal effect "
    "of customary coffee consumption on AF, and with high certainty that the residual "
    "signal reflects adiposity-linked pleiotropy rather than caffeine pharmacology; "
    "small effects in either direction cannot be excluded, and routine caffeine "
    "restriction, including after cardioversion, remains unsupported.")
paras = d.paragraphs
ai = next(i for i, p in enumerate(paras) if p.text.strip() == "Abstract")
ri = next(i for i, p in enumerate(paras) if p.text.strip() == "References")
bw = sum(len(p.text.split()) for p in paras[ai:ri])
aw = sum(len(p.text.split()) for p in paras)
wc_p = next(p for p in paras if p.text.strip().startswith("Word count:"))
set_text(wc_p, re.sub(r"Word count: [\d,]+ \(abstract through conclusions\); [\d,]+ all-inclusive",
                      f"Word count: {bw:,} (abstract through conclusions); {aw:,} all-inclusive",
                      wc_p.text))
d.save(f"{BASE}/咖啡与房颤_投稿版_v2.docx")
print(f"EN conclusions updated; words {bw:,}/{aw:,}")

dc = Document(f"{BASE}/论文中文审阅版.docx")
replace_in_doc(dc,
    "客观生物标志物 MR 为零；且未见真实世界安全信号。",
    "客观生物标志物 MR 为零；遗传预测的摄入并不升高循环咖啡因家族代谢物；且未见真实世界安全信号。")
replace_in_doc(dc,
    "现有证据反对习惯性咖啡饮用对房颤存在中等到大的因果效应，但任一方向的小效应无法排除；"
    "常规限制咖啡因——包括复律后——仍缺乏证据。",
    "现有证据以中等确定性反对习惯性咖啡饮用对房颤存在中等到大的因果效应，并以高确定性表明残余信号"
    "反映肥胖相关多效性而非咖啡因药理；任一方向的小效应仍无法排除，常规限制咖啡因——包括复律后——"
    "缺乏证据支持。")
dc.save(f"{BASE}/论文中文审阅版.docx")
print("CN conclusions updated")
