# -*- coding: utf-8 -*-
"""Final-audit fixes (EN + CN): stale 'three' framing, abstract first sentence,
strengths sentence, S1/S4/S9 citations, extended data availability."""
import re
from docx import Document

BASE = "方案B_合并稿/投稿_HeartRhythm_压缩版"


def set_text(p, t):
    p.runs[0].text = t
    for r in p.runs[1:]:
        r.text = ""


def replace_in_doc(doc, old, new, must=True):
    for p in doc.paragraphs:
        if old in p.text:
            set_text(p, p.text.replace(old, new))
            return True
    if must:
        raise RuntimeError("not found: " + old[:60])
    return False


# ================= EN =================
d = Document(f"{BASE}/咖啡与房颤_投稿版_v2.docx")

replace_in_doc(d, "by triangulating total genetic effects, mediation pathways, and real-world safety.",
               "by triangulating genetic, mechanistic, and real-world evidence.")
replace_in_doc(d, "This study combines three components using publicly available data",
               "This study combines complementary designs using publicly available data")
replace_in_doc(d, "data sources are in Table S11.",
               "data sources are in Table S11 and indexed in Table S4.")
replace_in_doc(d, "Strengths include three orthogonally biased designs;",
               "Strengths include four design families with orthogonal biases;")
replace_in_doc(d, "Reporting follows STROBE-MR [14] and READUS-PV [15].",
               "Reporting follows STROBE-MR [14] and READUS-PV [15] (Supplementary Table S1).")
replace_in_doc(d, "Colocalization separated the outlier loci into two classes (Supplementary Table S16).",
               "Colocalization separated the outlier loci into two classes (Supplementary Tables S9, S16).")
replace_in_doc(d, "cross-ancestry meta-GWAS (GWAS Catalog GCST90204201). FAERS data",
               "cross-ancestry meta-GWAS (GWAS Catalog GCST90204201). Sensitivity and mechanism "
               "analyses additionally used the PAGE coffee-intake GWAS (OpenGWAS "
               "ebi-a-GCST008028), FinnGen R12 arrhythmia-spectrum endpoints, the Chen et al. "
               "plasma metabolome GWAS (GWAS Catalog GCST90199644\u2013GCST90200436), "
               "electrophysiology GWAS (QT interval GCST90165290; resting heart rate "
               "GCST007609; PR interval GCST010321), and Ensembl GRCh37 gene annotation "
               "(REST API). FAERS data")

# word count update
paras = d.paragraphs
ai = next(i for i, p in enumerate(paras) if p.text.strip() == "Abstract")
ri = next(i for i, p in enumerate(paras) if p.text.strip() == "References")
bw = sum(len(p.text.split()) for p in paras[ai:ri])
aw = sum(len(p.text.split()) for p in paras)
wc_p = next(p for p in paras if p.text.strip().startswith("Word count:"))
set_text(wc_p, re.sub(r"Word count: [\d,]+ \(abstract through conclusions\); [\d,]+ all-inclusive",
                      f"Word count: {bw:,} (abstract through conclusions); {aw:,} all-inclusive",
                      wc_p.text))
d.save(f"{BASE}/咖啡与房颤_投稿版_v2.docx")
print(f"EN fixed; words {bw:,}/{aw:,}")

# ================= CN =================
dc = Document(f"{BASE}/论文中文审阅版.docx")
replace_in_doc(dc, "我们通过三角验证总遗传效应、中介通路与真实世界安全性，重新审视咖啡与房颤的关系。",
               "我们通过三角验证遗传、机制与真实世界证据，重新审视咖啡与房颤的关系。")
replace_in_doc(dc, "本研究结合三个使用公开数据的组成部分",
               "本研究结合多个使用公开数据的互补设计")
replace_in_doc(dc, "数据来源见表 S11。", "数据来源见表 S11，文件索引见表 S4。")
replace_in_doc(dc, "优点包括三个偏倚正交的设计；", "优点包括四个偏倚正交的设计族；")
replace_in_doc(dc, "报告遵循 STROBE-MR[14] 与 READUS-PV[15]。",
               "报告遵循 STROBE-MR[14] 与 READUS-PV[15]（补充表 S1）。")
replace_in_doc(dc, "共定位将离群位点分为两类（补充表 S16）",
               "共定位将离群位点分为两类（补充表 S9、S16）")
replace_in_doc(dc, "跨族裔 Meta-GWAS（GWAS Catalog GCST90204201）。FAERS 数据经 openFDA API 获取",
               "跨族裔 Meta-GWAS（GWAS Catalog GCST90204201）。敏感性与机制分析另使用 PAGE 咖啡摄入 GWAS"
               "（OpenGWAS ebi-a-GCST008028）、FinnGen R12 心律失常谱系端点、Chen 等血浆代谢组 GWAS"
               "（GWAS Catalog GCST90199644–GCST90200436）、电生理 GWAS（QT 间期 GCST90165290；"
               "静息心率 GCST007609；PR 间期 GCST010321）及 Ensembl GRCh37 基因注释（REST API）。"
               "FAERS 数据经 openFDA API 获取")
dc.save(f"{BASE}/论文中文审阅版.docx")
print("CN fixed")
